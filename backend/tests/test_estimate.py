"""Смета по нормативам пакета. Приёмка: заявка Z-01 = смета 01_Смета_пример_Орда_Трейдинг.xlsx."""
import datetime as dt
from decimal import Decimal

import openpyxl
import pytest

from app.services import estimate as est
from app.services.norms import load_norms
from tests.conftest import DATA

AS_OF = dt.date(2026, 10, 5)  # дата сметы-примера
EXAMPLE = DATA.parent / "01_Смета_пример_Орда_Трейдинг.xlsx"
REQUESTS = DATA.parent / "10_Заявки_на_расчёт.xlsx"


@pytest.fixture(scope="module")
def norms():
    return load_norms()


def _requests():
    ws = openpyxl.load_workbook(REQUESTS, data_only=True)["Заявки"]
    header = [c.value for c in ws[1]]
    return [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]


def _run(norms, req, **kw):
    return est.estimate(
        norms,
        city=req["Город"],
        segment=req["Сегмент"],
        area_m2=Decimal(str(req["Площадь, м²"])),
        toilets=int(req["Санузлов"] or 0),
        base_schedule=req["График базовой уборки"],
        support=est.parse_support(req["Поддерживающая уборка: позиций в смене × график"]),
        hygiene=req["Гигиена по ТЗ"] == "да",
        general_cleaning=req["Генуборка"] == "да",
        budget_gross_month=Decimal(str(req["Бюджет клиента в месяц, тг с НДС"])) if req["Бюджет клиента в месяц, тг с НДС"] else None,
        as_of=AS_OF,
        **kw,
    )


def _example():
    """Пересчёт листов примера по их же формулам: количества и цены берутся из файла, не из норм."""
    ws = openpyxl.load_workbook(EXAMPLE, data_only=False)["Смета"]
    rows = {i: r for i, r in enumerate(ws.iter_rows(min_row=1, max_row=76, values_only=True), start=1)}
    # ФОТ: строки 28–30 — C × D; строка 31 — генуборка C31 × (16000 × 4 / 12); E32 — gross-up 6 %
    payroll = sum((Decimal(rows[i][2]) * Decimal(rows[i][3]) for i in (28, 29, 30)), Decimal(0))
    payroll += Decimal(rows[31][2]) * Decimal(16000) * 4 / 12
    fot = payroll + payroll * Decimal("0.06") / (1 - Decimal("0.06"))
    # материалы, гигиена, инвентарь, оборудование, спецодежда: H = D × (F / 1,16) / G
    items = {}
    for i in range(37, 77):
        r = rows[i]
        if isinstance(r[0], int) and isinstance(r[5], (int, float)):
            items[r[1]] = (int(r[3]), Decimal(r[3]) * (Decimal(r[5]) / Decimal("1.16")) / Decimal(r[6]))
    return fot, items


def test_norms_loaded(norms):
    assert norms.warnings == []  # 07 и 09 согласованы
    assert (norms.vat_on(dt.date(2025, 12, 31)), norms.vat_on(dt.date(2026, 1, 1))) == (Decimal("0.12"), Decimal("0.16"))
    assert (norms.ic_margin, norms.min_ic_margin, norms.platform) == (Decimal("0.15"), Decimal("0.1"), Decimal("0.06"))
    assert norms.productivity["Офисы и бизнес-центры"] == 48
    assert norms.schedules["2/2 день 12 ч"].people_per_position == 2
    assert norms.city_k("Шымкент") == Decimal("0.85")


def test_z01_matches_example_estimate(norms):
    z01 = next(r for r in _requests() if r["Заявка"] == "Z-01")
    r = _run(norms, z01)
    fot, items = _example()
    assert abs(r.totals["payroll"] - fot) < Decimal("0.01")
    ours = {l.name: l for l in r.lines if l.section != "payroll"}
    assert len(items) == 23 and set(items) == set(ours)
    for name, (qty, monthly) in items.items():
        line = ours[name]
        assert line.qty == qty, name
        assert abs(line.monthly_net - monthly) < Decimal("0.01"), name
    cost = fot + sum((v for _, v in items.values()), Decimal(0))
    fee = cost * Decimal("0.101") / Decimal("0.9")
    admin = (cost + fee) * Decimal("0.01") / Decimal("0.99")
    budget = cost + fee + admin
    price = budget + budget * 15 / 85
    assert abs(r.cost - cost) < Decimal("0.01")
    assert abs(r.subcontractor_budget - budget) < Decimal("0.01")
    assert abs(r.price_net - price) < Decimal("0.01")
    assert abs(r.price_gross - price * Decimal("1.16")) < Decimal("0.01")
    assert abs(r.budget_gap - (Decimal(37200000) / 12 - price * Decimal("1.16"))) < Decimal("0.01")
    assert (r.operators, r.base_positions) == (8, 6)


@pytest.mark.parametrize("request_id", ["Z-02", "Z-03", "Z-04"])
def test_other_requests_compute(norms, request_id):
    req = next(r for r in _requests() if r["Заявка"] == request_id)
    r = _run(norms, req)
    assert r.price_net > r.subcontractor_budget > r.cost > 0
    assert abs(r.ic_margin / r.price_net - Decimal("0.15")) < Decimal("0.0001")


def test_warehouse_equipment_and_shift_staffing(norms):
    z02 = next(r for r in _requests() if r["Заявка"] == "Z-02")
    r = _run(norms, z02)
    # 12 000 м² ÷ 450 м²/ч·ч ÷ 12 ч = 2,2 → 3 позиции × 2 чел. (2/2) + поддержка 1 × 2 чел.
    assert (r.base_positions, r.operators) == (3, 8)
    eq = {l.name: l.qty for l in r.lines if l.section == "equipment"}
    assert eq == {"Поломоечная машина с сиденьем": 2, "Пылеводосос": 1}
    assert not [l for l in r.lines if l.section == "hygiene"]  # гигиена не по ТЗ


def test_city_coefficient_for_regions(norms):
    z04 = next(r for r in _requests() if r["Заявка"] == "Z-04")
    r = _run(norms, z04)
    base = next(l for l in r.lines if l.name.startswith("ОПУ базовая"))
    assert base.price_net == Decimal(118000) * Decimal("0.85")
    assert any("Шымкент" in w for w in r.warnings)


def test_budget_below_minimum_margin(norms):
    z01 = next(r for r in _requests() if r["Заявка"] == "Z-01")
    r = _run(norms, {**z01, "Бюджет клиента в месяц, тг с НДС": 2_100_000})
    assert r.budget_gap < 0 and r.margin_at_budget_pct < 10
    assert any("согласия директора" in w for w in r.warnings)


def test_vat12_is_flagged(norms):
    z01 = next(r for r in _requests() if r["Заявка"] == "Z-01")
    r = _run(norms, z01, vat_rate=Decimal("0.12"))
    assert r.vat == est.q2(r.price_net * Decimal("0.12"))
    assert any("для сравнения" in w for w in r.warnings)


def test_errors(norms):
    z01 = next(r for r in _requests() if r["Заявка"] == "Z-01")
    with pytest.raises(est.EstimateError, match="сегмента"):
        _run(norms, {**z01, "Сегмент": "Космодром"})
    with pytest.raises(est.EstimateError, match="графика"):
        _run(norms, {**z01, "График базовой уборки": "7/0 круглосуточно"})
    with pytest.raises(est.EstimateError, match="N × график"):
        est.parse_support("пара человек вечером")
