"""Нормативы из пакета организаторов.

07_Нормативы.json     — финансовые параметры (вознаграждение и админрасходы субподрядчика, маржа IC,
                        платформа занятости, НДС по датам, почасовые ставки операторов).
09_Нормы_для_сметы.xlsx — производительность по сегментам, ставки ОПУ по графикам, коэффициенты городов,
                        расход и цены материалов, гигиены, инвентаря, оборудования, спецодежды, генуборки.

Файлы читаются как данные; любое расхождение между ними попадает в warnings, а не исправляется молча.
"""
import datetime as dt
import json
import re
from dataclasses import dataclass, field
from decimal import Decimal
from functools import lru_cache
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl

DATA_DIR = Path(__file__).resolve().parents[3] / "data"
FIN_FILE = "07_Нормативы.json"
ESTIMATE_FILE = "09_Нормы_для_сметы.xlsx"


class NormsError(Exception):
    pass


def D(v) -> Decimal:
    return Decimal(str(v))


@dataclass(frozen=True)
class Schedule:
    name: str  # «5/2 утро 4 ч»
    monthly_rate: Decimal  # ставка в месяц для Алматы/Астаны
    hours_per_shift: Decimal
    people_per_position: Decimal


@dataclass(frozen=True)
class Item:
    name: str
    unit: str
    qty_rule: Decimal  # расход на 1000 м² / на санузел / правило
    price_with_vat: Decimal
    life_months: int = 1
    rule: str = ""


@dataclass(frozen=True)
class EquipmentItem:
    segment: str
    name: str
    rule: str
    price_with_vat: Decimal
    life_months: int


@dataclass
class Norms:
    # 07
    vat_by_date: List[Tuple[dt.date, Decimal]]  # (с какой даты, ставка)
    subcontractor_fee: Decimal
    subcontractor_admin: Decimal
    ic_margin: Decimal
    platform: Decimal
    operator_rates_per_hour: Dict[str, Decimal]
    # 09
    productivity: Dict[str, Decimal]  # сегмент → м² на человеко-час
    material_coef: Dict[str, Decimal]  # сегмент → коэффициент расхода к офису
    schedules: Dict[str, Schedule]
    city_coef: Dict[str, Decimal]
    materials: List[Item]
    hygiene: List[Item]
    inventory: List[Item]
    equipment: List[EquipmentItem]
    uniform_sets_per_year: Decimal
    uniform_price: Decimal
    shoes_per_year: Decimal
    shoes_price: Decimal
    general_cleaning_pay: Decimal
    general_cleanings_per_year: Decimal
    min_ic_margin: Decimal
    sources: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)

    def vat_on(self, day: dt.date) -> Decimal:
        rate = None
        for start, r in sorted(self.vat_by_date):
            if day >= start:
                rate = r
        if rate is None:
            raise NormsError(f"Нет ставки НДС на {day}")
        return rate

    def city_k(self, city: str) -> Decimal:
        return self.city_coef.get(city, self.city_coef.get("другие города", Decimal("1")))


# ---------- 07_Нормативы.json ----------
_VAT_KEY = re.compile(r"(до|с)\s+(\d{4}-\d{2}-\d{2})")


def _parse_vat(raw: Dict[str, float]) -> List[Tuple[dt.date, Decimal]]:
    out = []
    for key, rate in raw.items():
        m = _VAT_KEY.search(key)
        if not m:
            raise NormsError(f"Не понял период НДС «{key}»")
        day = dt.date.fromisoformat(m.group(2))
        # «до 2025-12-31» — действует с начала времён; «с 2026-01-01» — с этой даты
        out.append((dt.date.min if m.group(1) == "до" else day, D(rate)))
    return out


def _load_financial(path: Path) -> dict:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as e:
        raise NormsError(f"Не удалось прочитать {path.name}: {e}") from e


# ---------- 09_Нормы_для_сметы.xlsx ----------
def _rows(ws) -> List[tuple]:
    return [r for r in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]


def _sheet(wb, name: str):
    if name not in wb.sheetnames:
        raise NormsError(f"В {ESTIMATE_FILE} нет листа «{name}»")
    return wb[name]


def _load_estimate(path: Path) -> dict:
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except (OSError, ValueError, KeyError) as e:
        raise NormsError(f"Не удалось прочитать {path.name}: {e}") from e

    out: dict = {}
    prod = _rows(_sheet(wb, "Производительность"))
    out["productivity"] = {r[0]: D(r[1]) for r in prod}
    out["material_coef"] = {r[0]: D(r[2]) for r in prod}
    out["schedules"] = {
        r[0]: Schedule(r[0], D(r[1]), D(r[2]), D(r[3])) for r in _rows(_sheet(wb, "Графики и ставки ОПУ"))
    }
    out["city_coef"] = {r[0]: D(r[1]) for r in _rows(_sheet(wb, "Коэффициенты городов"))}
    out["materials"] = [Item(r[0], r[1], D(r[2]), D(r[3])) for r in _rows(_sheet(wb, "Материалы"))]
    out["hygiene"] = [Item(r[0], r[1], D(r[2]), D(r[3])) for r in _rows(_sheet(wb, "Гигиена"))]
    out["inventory"] = [
        Item(r[0], "шт", Decimal(0), D(r[2]), int(r[3]), rule=r[1]) for r in _rows(_sheet(wb, "Инвентарь"))
    ]
    out["equipment"] = [
        EquipmentItem(r[0], r[1], r[2], D(r[3]), int(r[4])) for r in _rows(_sheet(wb, "Оборудование"))
    ]
    uniform = {r[0]: r[1] for r in _rows(_sheet(wb, "Спецодежда и генуборка"))}
    try:
        out["uniform_sets_per_year"] = D(uniform["Форма ОПУ, комплектов на 1 ОПУ в год"])
        out["uniform_price"] = D(uniform["Цена формы с НДС, тг"])
        out["shoes_per_year"] = D(uniform["Обувь, пар на 1 ОПУ в год"])
        out["shoes_price"] = D(uniform["Цена обуви с НДС, тг"])
        out["general_cleaning_pay"] = D(uniform["Генеральная уборка: оплата 1 ОПУ за одну уборку, тг"])
        out["general_cleanings_per_year"] = D(uniform["Генеральная уборка: раз в год"])
    except KeyError as e:
        raise NormsError(f"В листе «Спецодежда и генуборка» нет строки {e}") from e
    fin = {r[0]: r[1] for r in _rows(_sheet(wb, "Финансовые параметры"))}
    out["sheet_fin"] = fin
    wb.close()
    return out


@lru_cache(maxsize=4)
def load_norms(data_dir: Optional[str] = None) -> Norms:
    base = Path(data_dir) if data_dir else DATA_DIR
    fin_path, est_path = base / FIN_FILE, base / ESTIMATE_FILE
    for p in (fin_path, est_path):
        if not p.exists():
            raise NormsError(f"Нет файла нормативов {p}")
    fin = _load_financial(fin_path)
    est = _load_estimate(est_path)
    sheet_fin = est.pop("sheet_fin")

    try:
        norms = Norms(
            vat_by_date=_parse_vat(fin["vat_rate"]),
            subcontractor_fee=D(fin["subcontractor_fee_share_of_budget"]),
            subcontractor_admin=D(fin["subcontractor_admin_share_of_budget"]),
            ic_margin=D(fin["ic_margin_share_of_price"]),
            platform=D(fin["platform_gross_up"]),
            operator_rates_per_hour={k: D(v) for k, v in fin["operator_rates_tg_per_hour"].items()},
            min_ic_margin=D(sheet_fin.get("Минимально допустимая маржа IC Line", 0)),
            sources=[FIN_FILE, ESTIMATE_FILE],
            **est,
        )
    except KeyError as e:
        raise NormsError(f"В {FIN_FILE} нет параметра {e}") from e

    # Сверка двух источников финансовых параметров
    pairs = [
        ("Расходы Платформы Занятости", norms.platform),
        ("Вознаграждение субподрядчика", norms.subcontractor_fee),
        ("Административные расходы субподрядчика", norms.subcontractor_admin),
        ("Маржа IC Line", norms.ic_margin),
    ]
    for label, value in pairs:
        if label in sheet_fin and D(sheet_fin[label]) != value:
            norms.warnings.append(f"«{label}»: в {ESTIMATE_FILE} {sheet_fin[label]}, в {FIN_FILE} {value} — взято из {FIN_FILE}")
    return norms
