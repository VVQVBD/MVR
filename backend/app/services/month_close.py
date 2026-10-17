"""«Закрытие месяца» для экрана frontend/app/page.tsx — считает бекенд, не статический JSON.

Выручка, выплаты субподрядчику и источники берутся из уже загруженных в БД договоров, объектов
(их `sources`) и актов субподрядчиков (их строки уже сопоставлены объектам при импорте —
см. app/services/actuals.py). Два поля бекенд не хранит как сущности и поэтому читает
на лету прямо из файла фактов: замечания клиентов и число «непривязанных» отметок смен
(сам файл графика, 04_Операторы_и_график_сентябрь.xlsx, организаторами не предоставлен).
"""
import datetime as dt
import re
from decimal import Decimal
from pathlib import Path
from typing import Dict, List, Optional

import anyio
import openpyxl
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app import models as m

DATA_DIR = Path(__file__).resolve().parents[3] / "data"
ACTUALS_FILE = DATA_DIR / "05_Факт_сентябрь_2026.xlsx"

MALLS = [
    {"id": "nord", "title": "ТРЦ «Норд»", "name_pattern": re.compile("Норд", re.I)},
    {"id": "samal-city", "title": "ТРЦ «Самал Сити»", "name_pattern": re.compile("Самал\\s*Сити", re.I)},
]
# Клиентский фронтенд сам раскладывает общий реестр disputes по объектам этим же правилом —
# backend отдаёт полный список, как /api/v1/contracts/discrepancies.
SOURCE_POLICY = (
    "Финансовые значения хранят ссылку на исходный документ. "
    "Смены без объектной связи не распределяются автоматически."
)
_ROW_RE = re.compile(r"строка\s+(\d+)")


def source(file: str, sheet: Optional[str] = None, row: Optional[int] = None, field: Optional[str] = None) -> dict:
    out: dict = {"file": file}
    if sheet:
        out["sheet"] = sheet
    if row is not None:
        out["row"] = row
    if field:
        out["field"] = field
    return out


def audited(value, sources: List[dict], calculation: Optional[str] = None) -> dict:
    out = {"value": value, "sources": sources}
    if calculation:
        out["calculation"] = calculation
    return out


def object_sources(obj: m.FacilityObject) -> List[dict]:
    return [source(s.file, s.sheet, s.row, None) for s in obj.sources] or [
        source(s.file, None, None, None) for s in (obj.contract.sources if obj.contract else [])
    ]


def _line_row(line: m.BillingActLine) -> Optional[int]:
    match = _ROW_RE.search(line.note or "")
    return int(match.group(1)) if match else None


def payout_sources(lines: List[m.BillingActLine]) -> List[dict]:
    out = []
    for line in lines:
        act = line.act
        file = act.source.file if act.source else ACTUALS_FILE.name
        sheet = act.source.sheet if act.source else None
        out.append(source(file, sheet, _line_row(line), "Сумма, тг (без НДС)"))
    return out


def _mall_of(text: Optional[str]) -> Optional[str]:
    low = text or ""
    for mall in MALLS:
        if mall["name_pattern"].search(low):
            return mall["id"]
    return None


def _reconciliation_row(obj: m.FacilityObject, lines_by_object: Dict[int, List[m.BillingActLine]]) -> dict:
    lines = lines_by_object.get(obj.id, [])
    revenue = Decimal(obj.monthly_tariff or 0)
    payout = sum((l.amount_net for l in lines), Decimal(0))
    rev_src = object_sources(obj)
    pay_src = payout_sources(lines)
    anomalies = []
    if not lines:
        anomalies.append(
            {
                "code": "MISSING_SUBCONTRACTOR_ACT",
                "severity": "medium",
                "message": "Для договорного объекта не найдена строка в акте субподрядчика.",
                "sources": rev_src,
            }
        )
    return {
        "id": f"{_mall_of(obj.name)}:{obj.contract.number}",
        "client": obj.client_company.name,
        "zone": obj.name,
        "subcontractor": obj.subcontractor.name if obj.subcontractor else None,
        "vatBasis": obj.vat_note or "не указан",
        "contract": audited(obj.contract.number, rev_src),
        "revenue": audited(revenue, rev_src, "Договорный месячный тариф"),
        "subcontractorPayout": audited(payout, pay_src, "Сумма строк акта субподрядчика"),
        "operatorPayout": audited(None, [], "Не рассчитывается: отсутствует связь отметок с объектом и планом смен"),
        "marginBeforeOperatorPayout": audited(
            revenue - payout, rev_src + pay_src, "Выручка − выплата субподрядчику; перед использованием проверить базу НДС"
        ),
        "anomalies": anomalies,
        "status": "review" if anomalies else "closed",
    }


def _rows(ws) -> List[tuple]:
    header = [str(c.value) if c.value is not None else "" for c in ws[1]]
    out = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        out.append((i, dict(zip(header, row))))
    return out


def _client_remarks_sync(mall_id: str) -> List[dict]:
    if not ACTUALS_FILE.exists():
        return []
    wb = openpyxl.load_workbook(ACTUALS_FILE, data_only=True, read_only=True)
    if "Замечания клиентов" not in wb.sheetnames:
        wb.close()
        return []
    out = []
    for row_num, row in _rows(wb["Замечания клиентов"]):
        if _mall_of(row.get("Объект")) != mall_id:
            continue
        resolved = row.get("Устранено")
        out.append(
            {
                "id": str(row.get("ID") or ""),
                "description": str(row.get("Описание") or ""),
                "detectedAt": str(row.get("Выявлено") or ""),
                "dueAt": str(row.get("Срок устранения") or ""),
                "resolvedAt": str(resolved) if resolved else None,
                "status": "resolved" if resolved else "open",
                "sources": [source(ACTUALS_FILE.name, "Замечания клиентов", row_num)],
            }
        )
    wb.close()
    return out


def _unlinked_shifts_count() -> int:
    if not ACTUALS_FILE.exists():
        return 0
    wb = openpyxl.load_workbook(ACTUALS_FILE, data_only=True, read_only=True)
    if "Отметки" not in wb.sheetnames:
        wb.close()
        return 0
    count = sum(1 for row in wb["Отметки"].iter_rows(min_row=2, values_only=True) if any(v is not None for v in row))
    wb.close()
    return count


async def _client_remarks_async(mall_id: str) -> List[dict]:
    return await anyio.to_thread.run_sync(_client_remarks_sync, mall_id)


async def _unlinked_shifts_count_async() -> int:
    return await anyio.to_thread.run_sync(_unlinked_shifts_count)


async def build(db: AsyncSession, period: dt.date) -> dict:
    objects = (
        await db.scalars(
            select(m.FacilityObject).options(
                selectinload(m.FacilityObject.sources),
                selectinload(m.FacilityObject.contract).selectinload(m.Contract.sources),
                selectinload(m.FacilityObject.client_company),
                selectinload(m.FacilityObject.subcontractor),
            )
        )
    ).all()
    lines = (
        await db.scalars(
            select(m.BillingActLine)
            .join(m.BillingAct)
            .where(m.BillingAct.kind == "subcontractor_act", m.BillingAct.period_month == period)
            .options(selectinload(m.BillingActLine.act).selectinload(m.BillingAct.source))
        )
    ).all()
    lines_by_object: Dict[int, List[m.BillingActLine]] = {}
    for line in lines:
        if line.facility_object_id is not None:
            lines_by_object.setdefault(line.facility_object_id, []).append(line)

    disputes_rows = (
        await db.scalars(select(m.DisputedCondition).options(selectinload(m.DisputedCondition.sources)))
    ).all()
    disputes = [
        {
            "id": d.id,
            "severity": d.severity,
            "topic": d.topic,
            "description": d.issue,
            "request": d.ask,
            "status": d.status,
            "sources": [source(s.file, s.sheet, s.row) for s in d.sources],
        }
        for d in disputes_rows
    ]

    out_objects = []
    for mall in MALLS:
        mall_objects = [o for o in objects if _mall_of(o.name) == mall["id"]]
        reconciliation = [_reconciliation_row(o, lines_by_object) for o in mall_objects]
        remarks = await _client_remarks_async(mall["id"])
        open_remarks = sum(1 for r in remarks if r["status"] == "open")
        kpi = {
            "revenue": sum((r["revenue"]["value"] for r in reconciliation), Decimal(0)),
            "subcontractorPayout": sum((r["subcontractorPayout"]["value"] for r in reconciliation), Decimal(0)),
            "marginBeforeOperatorPayout": sum(
                (r["marginBeforeOperatorPayout"]["value"] for r in reconciliation), Decimal(0)
            ),
            "discrepancies": open_remarks + sum(len(r["anomalies"]) for r in reconciliation),
            "operatorPayout": None,
            "operatorPayoutStatus": "unavailable_without_shift_object_mapping",
        }
        out_objects.append(
            {
                "id": mall["id"],
                "title": mall["title"],
                "reconciliation": reconciliation,
                "clientRemarks": remarks,
                "kpi": kpi,
            }
        )

    return {
        "schemaVersion": "1.0",
        "period": f"{period:%Y-%m}",
        "currency": "KZT",
        "sourcePolicy": SOURCE_POLICY,
        "objects": out_objects,
        "unlinkedShiftsCount": await _unlinked_shifts_count_async(),
        "disputes": disputes,
    }
